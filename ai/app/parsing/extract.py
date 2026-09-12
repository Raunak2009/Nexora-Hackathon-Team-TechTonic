"""Raw text extraction from resume / JD files.

Handles PDF (the format the judges will hand us), DOCX, legacy DOC, TXT, MD and
RTF. Every extractor is best-effort and degrades gracefully: a file we cannot
read becomes an empty document with a recorded warning rather than a crash that
takes the whole batch down.
"""

from __future__ import annotations

import re
import shutil
import subprocess
import tempfile
import unicodedata
from dataclasses import dataclass, field
from pathlib import Path

from ..config import SUPPORTED_EXTENSIONS


@dataclass
class ExtractedDocument:
    path: Path
    text: str
    method: str
    warnings: list[str] = field(default_factory=list)

    @property
    def ok(self) -> bool:
        return len(self.text.strip()) >= 40


# ---------------------------------------------------------------------------
# Individual format handlers
# ---------------------------------------------------------------------------

def _extract_pdf(path: Path) -> tuple[str, str, list[str]]:
    warnings: list[str] = []

    # 1st choice: PyMuPDF - fastest and best at multi-column resume layouts.
    # `pymupdf` is the current import name; `fitz` is the legacy alias and warns
    # on newer versions, so try the new name first.
    try:
        try:
            import pymupdf as fitz  # type: ignore
        except ImportError:
            import fitz  # type: ignore

        with fitz.open(path) as doc:
            pages = [page.get_text("text") for page in doc]
        text = "\n".join(pages)
        if text.strip():
            return text, "pymupdf", warnings
        warnings.append("PyMuPDF returned no text (scanned/image PDF?)")
    except ImportError:
        pass
    except Exception as exc:  # pragma: no cover - defensive
        warnings.append(f"PyMuPDF failed: {exc}")

    # 2nd choice: pdfplumber - slower but very reliable on text PDFs.
    try:
        import pdfplumber  # type: ignore

        with pdfplumber.open(path) as pdf:
            pages = [p.extract_text() or "" for p in pdf.pages]
        text = "\n".join(pages)
        if text.strip():
            return text, "pdfplumber", warnings
        warnings.append("pdfplumber returned no text")
    except ImportError:
        pass
    except Exception as exc:  # pragma: no cover
        warnings.append(f"pdfplumber failed: {exc}")

    # 3rd choice: pypdf.
    try:
        from pypdf import PdfReader  # type: ignore

        reader = PdfReader(str(path))
        text = "\n".join((p.extract_text() or "") for p in reader.pages)
        if text.strip():
            return text, "pypdf", warnings
    except ImportError:
        pass
    except Exception as exc:  # pragma: no cover
        warnings.append(f"pypdf failed: {exc}")

    warnings.append(
        "No PDF text layer found. If this is a scanned resume it needs OCR "
        "(install pytesseract + tesseract) - it is being scored on an empty body."
    )
    return "", "none", warnings


def _extract_docx(path: Path) -> tuple[str, str, list[str]]:
    warnings: list[str] = []
    try:
        import docx  # type: ignore

        document = docx.Document(str(path))
        parts: list[str] = [p.text for p in document.paragraphs]
        # Resumes love putting the entire skills matrix inside a table.
        for table in document.tables:
            for row in table.rows:
                cells = [c.text.strip() for c in row.cells]
                # de-duplicate merged cells that repeat the same text
                deduped: list[str] = []
                for c in cells:
                    if c and (not deduped or deduped[-1] != c):
                        deduped.append(c)
                if deduped:
                    parts.append(" | ".join(deduped))
        return "\n".join(parts), "python-docx", warnings
    except ImportError:
        warnings.append("python-docx not installed")
    except Exception as exc:
        warnings.append(f"python-docx failed: {exc}")
    return "", "none", warnings


def _extract_doc(path: Path) -> tuple[str, str, list[str]]:
    """Legacy binary .doc (Word 97-2003)."""
    warnings: list[str] = []

    for tool, args in (
        ("antiword", ["antiword", str(path)]),
        ("catdoc", ["catdoc", str(path)]),
    ):
        if shutil.which(tool):
            try:
                out = subprocess.run(args, capture_output=True, timeout=60)
                text = out.stdout.decode("utf-8", errors="replace")
                if text.strip():
                    return text, tool, warnings
            except Exception as exc:
                warnings.append(f"{tool} failed: {exc}")

    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if soffice:
        try:
            with tempfile.TemporaryDirectory() as tmp:
                subprocess.run(
                    [soffice, "--headless", "--convert-to",
                     "txt:Text (encoded):UTF8", "--outdir", tmp, str(path)],
                    capture_output=True, timeout=180,
                )
                produced = list(Path(tmp).glob("*.txt"))
                if produced:
                    return produced[0].read_text("utf-8", errors="replace"), "libreoffice", warnings
        except Exception as exc:
            warnings.append(f"libreoffice failed: {exc}")

    # Last resort: pull printable ASCII runs straight out of the binary. Ugly,
    # but a partially-read resume ranks better than a silently dropped one.
    try:
        blob = path.read_bytes()
        chunks = re.findall(rb"[\x20-\x7e\r\n\t]{6,}", blob)
        text = "\n".join(c.decode("ascii", errors="ignore") for c in chunks)
        text = re.sub(r"\b(HYPERLINK|PAGEREF|MERGEFORMAT|Times New Roman|Calibri|"
                      r"Arial|Symbol|Wingdings|Normal\.dotm?|Microsoft Word)\b", " ", text)
        if len(text.strip()) > 200:
            warnings.append("Read .doc via binary string salvage - text may be noisy. "
                            "Install antiword for clean extraction.")
            return text, "binary-salvage", warnings
    except Exception as exc:  # pragma: no cover
        warnings.append(f"binary salvage failed: {exc}")

    warnings.append("Could not read legacy .doc file.")
    return "", "none", warnings


# Resume XML: <resume><applicant/><summary/><technicalSkills/>...
_XML_SECTION_HEADERS = {
    "summary": "PROFESSIONAL SUMMARY",
    "objective": "PROFESSIONAL SUMMARY",
    "education": "EDUCATION",
    "technicalskills": "TECHNICAL SKILLS",
    "skills": "TECHNICAL SKILLS",
    "experience": "EXPERIENCE",
    "workexperience": "EXPERIENCE",
    "internships": "EXPERIENCE",
    "projects": "PROJECTS",
    "certifications": "CERTIFICATIONS",
    "achievements": "ACHIEVEMENTS",
    "awards": "ACHIEVEMENTS",
    "extracurricular": "ACHIEVEMENTS",
    "activities": "ACHIEVEMENTS",
    "publications": "ACHIEVEMENTS",
}


def _extract_xml(path: Path) -> tuple[str, str, list[str]]:
    """Structured resume XML -> text that still has its sections.

    A generic XML-to-text dump would throw away the one thing that makes this
    format better than a PDF: the sections are already labelled. So instead of
    flattening, each known top-level element is emitted under a header the
    section detector recognises, and the skills block keeps its
    "Languages: Python, Bash" shape rather than becoming one long line.

    Falls back to a plain tag-strip for any XML that is not this schema, so an
    unexpected layout degrades instead of failing.
    """
    import xml.etree.ElementTree as ET

    warnings: list[str] = []
    try:
        root = ET.parse(path).getroot()
    except ET.ParseError as exc:
        # Malformed XML still often has readable text in it.
        raw = path.read_text("utf-8", errors="replace")
        stripped = re.sub(r"<[^>]+>", "\n", raw)
        warnings.append(f"XML did not parse ({exc}); fell back to tag stripping.")
        return stripped, "xml-salvage", warnings

    def text_of(el) -> str:
        return " ".join((el.text or "").split()) if el is not None else ""

    out: list[str] = []

    applicant = root.find("applicant")
    if applicant is not None:
        name = text_of(applicant.find("name"))
        if name:
            out.append(name)
        contact = [text_of(applicant.find(t))
                   for t in ("phone", "email", "linkedin", "github", "location")]
        contact = [c for c in contact if c]
        if contact:
            out.append(" | ".join(contact))
        role = text_of(applicant.find("targetRole"))
        if role:
            out.append(f"Target Role: {role}")
        out.append("")

    for child in root:
        tag = child.tag.lower()
        if tag == "applicant":
            continue

        header = _XML_SECTION_HEADERS.get(tag)
        out.append(header or tag.upper())

        if tag in ("summary", "objective") or (child.text and len(child) == 0):
            out.append(text_of(child))

        elif tag in ("technicalskills", "skills"):
            for group in child:
                label = group.get("category") or group.get("name") or ""
                value = text_of(group)
                out.append(f"{label}: {value}" if label else value)

        else:
            # Repeating blocks: experience / projects / certifications / ...
            blocks = list(child) or [child]
            for block in blocks:
                if block.text and len(block) == 0:
                    out.append(text_of(block))
                    continue
                title = text_of(block.find("title")) or text_of(block.find("degree"))
                dates = text_of(block.find("dates")) or text_of(block.find("date"))
                line = " | ".join(p for p in (title, dates) if p)
                if line:
                    out.append(line)
                for extra in ("school", "cgpa", "coursework", "stack", "issuer"):
                    value = text_of(block.find(extra))
                    if value:
                        out.append(value)
                for bullet in block.iter("bullet"):
                    value = text_of(bullet)
                    if value:
                        out.append(f"- {value}")
                if not line and not list(block.iter("bullet")):
                    leftover = " ".join(t.strip() for t in block.itertext() if t.strip())
                    if leftover:
                        out.append(leftover)
        out.append("")

    text = "\n".join(out)
    if len(text.strip()) < 40:
        leftover = "\n".join(t.strip() for t in root.itertext() if t.strip())
        warnings.append("Unrecognised XML schema; used a generic text dump.")
        return leftover, "xml-generic", warnings
    return text, "xml", warnings


def _extract_rtf(path: Path) -> tuple[str, str, list[str]]:
    raw = path.read_text("utf-8", errors="replace")
    try:
        from striprtf.striprtf import rtf_to_text  # type: ignore

        return rtf_to_text(raw), "striprtf", []
    except ImportError:
        pass
    text = re.sub(r"\\[a-z]+-?\d* ?", " ", raw)
    text = text.replace("{", " ").replace("}", " ")
    return text, "regex-rtf", ["striprtf not installed - used a crude RTF strip"]


# ---------------------------------------------------------------------------
# Normalisation
# ---------------------------------------------------------------------------

_BULLETS = "•▪◦‣·∙●○–—"
_LIGATURES = {
    "\ufb00": "ff", "\ufb01": "fi", "\ufb02": "fl", "\ufb03": "ffi", "\ufb04": "ffl",
    "\u2019": "'", "\u2018": "'", "\u201c": '"', "\u201d": '"',
    "\u2013": "-", "\u2014": "-", "\u00a0": " ", "\ufeff": "",
}


def _rejoin_wrapped(text: str) -> str:
    """Join a line onto the previous one only when it is clearly a wrap."""
    lines = text.split("\n")
    out: list[str] = []
    for line in lines:
        prev = out[-1] if out else ""
        wrapped = (
            len(prev) >= 45                     # long enough to be a wrapped line
            and (prev[-1:].islower() or prev[-1:] == ",")
            and line[:1].islower()
            and line.strip()
        )
        if wrapped and out:
            out[-1] = prev + " " + line
        else:
            out.append(line)
    return "\n".join(out)


def normalize_text(text: str) -> str:
    """Make messy resume text uniform without destroying its structure."""
    if not text:
        return ""

    for bad, good in _LIGATURES.items():
        text = text.replace(bad, good)
    text = unicodedata.normalize("NFKC", text)

    lines: list[str] = []
    for raw_line in text.splitlines():
        line = raw_line.replace("\t", " ")
        line = "".join(" " if ch in _BULLETS else ch for ch in line)
        # collapse "S K I L L S" style letter-spaced headers from PDF exports
        if len(line) > 6 and re.fullmatch(r"(?:[A-Za-z] ){3,}[A-Za-z]\s*", line):
            line = line.replace(" ", "")
        line = re.sub(r"[ ]{2,}", " ", line).strip()
        lines.append(line)

    out = "\n".join(lines)
    out = re.sub(r"\n{3,}", "\n\n", out)
    # Some PDFs emit one word per line; re-join obvious fragments. Only join
    # when the PREVIOUS line is long enough to be a wrapped sentence - the
    # unguarded version glued "Aarav Menon" onto the contact line below it and
    # the name was lost for every resume with a lowercase-ending name.
    out = _rejoin_wrapped(out)
    return out.strip()


# ---------------------------------------------------------------------------
# Public entry points
# ---------------------------------------------------------------------------

def extract_document(path: str | Path) -> ExtractedDocument:
    path = Path(path)
    suffix = path.suffix.lower()

    if not path.exists():
        return ExtractedDocument(path, "", "none", [f"File not found: {path}"])

    if suffix == ".pdf":
        text, method, warnings = _extract_pdf(path)
    elif suffix == ".docx":
        text, method, warnings = _extract_docx(path)
    elif suffix == ".doc":
        text, method, warnings = _extract_doc(path)
    elif suffix == ".xml":
        text, method, warnings = _extract_xml(path)
    elif suffix == ".rtf":
        text, method, warnings = _extract_rtf(path)
    elif suffix in {".txt", ".md"}:
        text, method, warnings = path.read_text("utf-8", errors="replace"), "plaintext", []
    else:
        return ExtractedDocument(
            path, "", "none",
            [f"Unsupported extension '{suffix}'. Supported: {sorted(SUPPORTED_EXTENSIONS)}"],
        )

    doc = ExtractedDocument(path, normalize_text(text), method, warnings)
    if not doc.ok and not doc.warnings:
        doc.warnings.append("Extracted text was suspiciously short.")
    return doc


def extract_folder(folder: str | Path) -> list[ExtractedDocument]:
    folder = Path(folder)
    if not folder.is_dir():
        raise NotADirectoryError(f"{folder} is not a directory")
    files = sorted(
        p for p in folder.rglob("*")
        if p.is_file() and p.suffix.lower() in SUPPORTED_EXTENSIONS and not p.name.startswith(("~$", "."))
    )
    return [extract_document(p) for p in files]
